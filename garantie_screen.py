# ============================================================
# garantie_screen.py
# Module Garanties - Prep@udit Desktop
# Export automatique de tous les onglets
# ============================================================

import os
import pandas as pd

from PyQt6.QtWidgets import (
    QApplication,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QFileDialog,
    QMessageBox,
    QHeaderView,
    QFrame,
    QTabWidget,
    QAbstractItemView,
    QInputDialog,
    QDialog,
    QProgressBar,
    QTableView,
)

from PyQt6.QtCore import (
    Qt,
    QAbstractTableModel,
    QModelIndex,
)

from PyQt6.QtGui import QFont


# ============================================================
# COULEURS
# ============================================================

SALMON_BG = "#f6c9a8"
HEADER_BG = "#c0392b"
DARK = "#2d2d2d"
BORDER = "#999999"
WHITE = "#FFFFFF"


# ============================================================
# DOSSIER D'EXPORT PREP@UDIT
# ============================================================

DOSSIER_EXPORT = r"C:\Users\HP\Desktop\prep@udit\EXTRACTION"


# ============================================================
# FENÊTRE DE PROGRESSION
# ============================================================

class ProgressDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setFixedSize(450, 160)

        self.setWindowTitle(
            "Traitement en cours"
        )

        self.setWindowModality(
            Qt.WindowModality.ApplicationModal
        )

        self.setStyleSheet(
            """
            QDialog {
                background-color: white;
                border: 2px solid #8C4318;
                border-radius: 8px;
            }

            QLabel {
                background: transparent;
                color: #2d2d2d;
            }

            QProgressBar {
                border: 1px solid #999999;
                border-radius: 7px;
                background-color: #eeeeee;
                height: 32px;
                text-align: center;
                font-size: 14px;
                font-weight: bold;
            }

            QProgressBar::chunk {
                background-color: #c0392b;
                border-radius: 6px;
            }
            """
        )

        layout = QVBoxLayout(self)

        layout.setContentsMargins(
            35,
            25,
            35,
            25
        )

        layout.setSpacing(12)

        self.title_label = QLabel(
            "Traitement des garanties en cours..."
        )

        self.title_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.title_label.setFont(
            QFont(
                "Arial",
                15,
                QFont.Weight.Bold
            )
        )

        self.progress_bar = QProgressBar()

        self.progress_bar.setRange(
            0,
            100
        )

        self.progress_bar.setValue(
            0
        )

        self.progress_bar.setTextVisible(
            True
        )

        self.lines_label = QLabel(
            "> Lignes traitées : 0 / 0"
        )

        self.lines_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.lines_label.setFont(
            QFont(
                "Arial",
                11
            )
        )

        self.wait_label = QLabel(
            "Veuillez patienter..."
        )

        self.wait_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.wait_label.setFont(
            QFont(
                "Arial",
                11,
                QFont.Weight.Bold
            )
        )

        self.wait_label.setStyleSheet(
            """
            color: #555555;
            """
        )

        layout.addWidget(
            self.title_label
        )

        layout.addWidget(
            self.progress_bar
        )

        layout.addWidget(
            self.lines_label
        )

        layout.addWidget(
            self.wait_label
        )

    # ========================================================
    # UPDATE
    # ========================================================

    def update_progress(
        self,
        percent,
        processed=0,
        total=0
    ):

        percent = max(
            0,
            min(
                100,
                int(percent)
            )
        )

        self.progress_bar.setValue(
            percent
        )

        if total:

            self.lines_label.setText(
                f"> Lignes traitées : "
                f"{processed:,} / {total:,}".replace(
                    ",",
                    " "
                )
            )

        else:

            self.lines_label.setText(
                f"> Lignes traitées : "
                f"{processed:,}".replace(
                    ",",
                    " "
                )
            )

        QApplication.processEvents()


# ============================================================
# MODELE TABLEAU
# ============================================================

class DataFrameModel(QAbstractTableModel):

    def __init__(
        self,
        dataframe=None,
        parent=None
    ):

        super().__init__(parent)

        if dataframe is None:
            self._df = pd.DataFrame()
        else:
            self._df = dataframe

    # ========================================================
    # NOMBRE DE LIGNES
    # ========================================================

    def rowCount(
        self,
        parent=QModelIndex()
    ):

        if parent.isValid():
            return 0

        return len(self._df)

    # ========================================================
    # NOMBRE DE COLONNES
    # ========================================================

    def columnCount(
        self,
        parent=QModelIndex()
    ):

        if parent.isValid():
            return 0

        return len(self._df.columns)

    # ========================================================
    # DONNEES
    # ========================================================

    def data(
        self,
        index,
        role=Qt.ItemDataRole.DisplayRole
    ):

        if not index.isValid():
            return None

        row = index.row()
        column = index.column()

        if (
            role == Qt.ItemDataRole.DisplayRole
            or role == Qt.ItemDataRole.EditRole
        ):

            try:

                valeur = self._df.iloc[
                    row,
                    column
                ]

                if valeur is None:
                    return ""

                try:

                    if pd.isna(valeur):
                        return ""

                except Exception:
                    pass

                return str(valeur)

            except Exception:
                return ""

        if role == Qt.ItemDataRole.TextAlignmentRole:

            nom_colonne = str(
                self._df.columns[column]
            ).lower()

            if (
                "montant" in nom_colonne
                or nom_colonne == "id"
                or "code agence" in nom_colonne
            ):

                return (
                    Qt.AlignmentFlag.AlignRight
                    | Qt.AlignmentFlag.AlignVCenter
                )

            return (
                Qt.AlignmentFlag.AlignLeft
                | Qt.AlignmentFlag.AlignVCenter
            )

        return None

    # ========================================================
    # EN-TETES
    # ========================================================

    def headerData(
        self,
        section,
        orientation,
        role=Qt.ItemDataRole.DisplayRole
    ):

        if role != Qt.ItemDataRole.DisplayRole:
            return None

        if orientation == Qt.Orientation.Horizontal:

            if 0 <= section < len(self._df.columns):

                return str(
                    self._df.columns[section]
                )

        else:

            return str(
                section + 1
            )

        return None

    # ========================================================
    # REMPLACER DATAFRAME
    # ========================================================

    def set_dataframe(
        self,
        dataframe
    ):

        self.beginResetModel()

        if dataframe is None:
            self._df = pd.DataFrame()
        else:
            self._df = dataframe

        self.endResetModel()


# ============================================================
# GARANTIE SCREEN
# ============================================================

class GarantieScreen(QWidget):

    def __init__(
        self,
        on_back=None
    ):

        super().__init__()

        self.on_back = on_back

        self.df = pd.DataFrame()

        self.progress_dialog = None

        self.fichier_source = None

        self.base_model = DataFrameModel()

        self.constituee_model = DataFrameModel()

        self.incomplete_model = DataFrameModel()

        self.retenue_model = DataFrameModel()

        self.echue_model = DataFrameModel()

        self.setStyleSheet(
            f"""
            QWidget {{
                background-color: {SALMON_BG};
                color: {DARK};
            }}

            QTabWidget::pane {{
                border: 1px solid {BORDER};
                background: white;
            }}

            QTabBar::tab {{
                background: #ead6c8;
                padding: 9px 18px;
                border: 1px solid {BORDER};
                font-weight: bold;
            }}

            QTabBar::tab:selected {{
                background: white;
            }}
            """
        )

        self._build_ui()

        self.tabs.tabBar().setTabTextColor(
            0,
            Qt.GlobalColor.darkGreen
        )

        self.tabs.tabBar().setTabTextColor(
            1,
            Qt.GlobalColor.darkYellow
        )

        self.tabs.tabBar().setTabTextColor(
            2,
            Qt.GlobalColor.darkYellow
        )

        self.tabs.tabBar().setTabTextColor(
            3,
            Qt.GlobalColor.red
        )

        self.tabs.tabBar().setTabTextColor(
            4,
            Qt.GlobalColor.blue
        )

    # ========================================================
    # INTERFACE
    # ========================================================

    def _build_ui(self):

        main_layout = QVBoxLayout(self)

        main_layout.setContentsMargins(
            20,
            15,
            20,
            20
        )

        main_layout.setSpacing(12)

        # ====================================================
        # HEADER
        # ====================================================

        header = QFrame()

        header.setStyleSheet(
            """
            QFrame {
                background-color: white;
                border: 2px solid #8C4318;
                border-radius: 8px;
            }
            """
        )

        header_layout = QHBoxLayout(header)

        header_layout.setContentsMargins(
            15,
            10,
            15,
            10
        )

        # ====================================================
        # RETOUR
        # ====================================================

        self.btn_back = QPushButton(
            "← Retour au Menu"
        )

        self.btn_back.setFixedSize(
            142,
            38
        )

        self.btn_back.setCursor(
            Qt.CursorShape.PointingHandCursor
        )

        self.btn_back.setStyleSheet(
            """
            QPushButton {
                background-color: #E67E00;
                color: white;
                border: none;
                border-radius: 6px;
                padding: 5px 12px;
                font-weight: bold;
                font-size: 12px;
            }

            QPushButton:hover {
                background-color: #D66F00;
            }

            QPushButton:pressed {
                background-color: #BF6000;
            }
            """
        )

        self.btn_back.clicked.connect(
            self._go_back
        )

        # ====================================================
        # TITRE
        # ====================================================

        title = QLabel(
            "GARANTIES"
        )

        title.setFont(
            QFont(
                "Georgia",
                28,
                QFont.Weight.Bold
            )
        )

        title.setStyleSheet(
            f"""
            color: {HEADER_BG};
            border: none;
            background: transparent;
            """
        )

        title.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        right_space = QLabel()

        right_space.setFixedWidth(
            142
        )

        right_space.setStyleSheet(
            """
            background: transparent;
            border: none;
            """
        )

        header_layout.addWidget(
            self.btn_back
        )

        header_layout.addStretch()

        header_layout.addWidget(
            title
        )

        header_layout.addStretch()

        header_layout.addWidget(
            right_space
        )

        main_layout.addWidget(
            header
        )

        # ====================================================
        # TABS
        # ====================================================

        self.tabs = QTabWidget()

        self.base_table = self._create_table(
            self.base_model
        )

        self.constituee_table = self._create_table(
            self.constituee_model
        )

        self.incomplete_table = self._create_table(
            self.incomplete_model
        )

        self.retenue_table = self._create_table(
            self.retenue_model
        )

        self.echue_table = self._create_table(
            self.echue_model
        )

        self.tabs.addTab(
            self.base_table,
            "Base Garanties"
        )

        self.tabs.addTab(
            self.constituee_table,
            "Constituée"
        )

        self.tabs.addTab(
            self.incomplete_table,
            "Incomplète"
        )

        self.tabs.addTab(
            self.retenue_table,
            "Retenue NC"
        )

        self.tabs.addTab(
            self.echue_table,
            "Echue"
        )

        main_layout.addWidget(
            self.tabs,
            1
        )

        # ====================================================
        # MESSAGE
        # ====================================================

        self.info_label = QLabel(
            "Sélectionnez un fichier Excel ou CSV pour lancer "
            "l'extraction des Garanties."
        )

        self.info_label.setStyleSheet(
            """
            color: #6b5b53;
            font-style: italic;
            padding: 4px;
            """
        )

        main_layout.addWidget(
            self.info_label
        )

    # ========================================================
    # CREATION TABLE
    # ========================================================

    def _create_table(
        self,
        model
    ):

        table = QTableView()

        table.setModel(
            model
        )

        table.setEditTriggers(
            QAbstractItemView.EditTrigger.NoEditTriggers
        )

        table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows
        )

        table.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )

        table.setAlternatingRowColors(
            True
        )

        table.setSortingEnabled(
            False
        )

        table.setWordWrap(
            False
        )

        table.setStyleSheet(
            f"""
            QTableView {{
                background-color: white;
                border: 1px solid {BORDER};
                gridline-color: #CCCCCC;
                alternate-background-color: #FAFAFA;
                selection-background-color: #FCE5CD;
                selection-color: black;
            }}

            QHeaderView::section {{
                background-color: {HEADER_BG};
                color: white;
                padding: 7px;
                border: 1px solid #999999;
                font-weight: bold;
            }}
            """
        )

        header = table.horizontalHeader()

        header.setSectionResizeMode(
            QHeaderView.ResizeMode.Interactive
        )

        header.setStretchLastSection(
            False
        )

        table.verticalHeader().setVisible(
            False
        )

        table.verticalHeader().setDefaultSectionSize(
            25
        )

        return table

    # ========================================================
    # PROGRESSION
    # ========================================================

    def _show_progress(self):

        self.progress_dialog = ProgressDialog(
            self
        )

        self.progress_dialog.show()

        QApplication.processEvents()

    def _set_progress(
        self,
        percent,
        processed=0,
        total=0
    ):

        if self.progress_dialog is not None:

            self.progress_dialog.update_progress(
                percent,
                processed,
                total
            )

    def _close_progress(self):

        if self.progress_dialog is not None:

            self.progress_dialog.close()

            self.progress_dialog.deleteLater()

            self.progress_dialog = None

            QApplication.processEvents()

    # ========================================================
    # RETOUR
    # ========================================================

    def _go_back(self):

        if self.on_back:
            self.on_back()

    # ========================================================
    # IMPORT
    # ========================================================

    def import_csv(self):

        fichier, _ = QFileDialog.getOpenFileName(
            self,
            "Sélectionner le fichier des Garanties",
            "",
            (
                "Fichiers Excel (*.xlsx *.xls *.xlsm);;"
                "Fichiers CSV (*.csv *.txt);;"
                "Tous les fichiers (*.*)"
            )
        )

        if not fichier:
            return False

        self.fichier_source = fichier

        self._show_progress()

        try:

            # =================================================
            # LECTURE
            # =================================================

            self._set_progress(
                10,
                0,
                0
            )

            extension = os.path.splitext(
                fichier
            )[1].lower()

            if extension in (
                ".xlsx",
                ".xls",
                ".xlsm"
            ):

                df = self._lire_excel(
                    fichier
                )

            else:

                df = self._lire_csv(
                    fichier
                )

            if df is None:

                self._close_progress()

                return False

            if df.empty:

                self._close_progress()

                QMessageBox.warning(
                    self,
                    "Garanties",
                    "Le fichier ne contient aucune donnée."
                )

                return False

            total = len(df)

            self._set_progress(
                25,
                0,
                total
            )

            self.df = self._preparer_base(
                df
            )

            self._set_progress(
                70,
                total,
                total
            )

            if self.df.empty:

                self._close_progress()

                QMessageBox.warning(
                    self,
                    "Garanties",
                    "Impossible de préparer les données."
                )

                return False

            # =================================================
            # AFFICHAGE
            # =================================================

            self._set_progress(
                80,
                total,
                total
            )

            self._afficher_resultats()

            # =================================================
            # EXPORT
            # =================================================

            self._set_progress(
                90,
                total,
                total
            )

            fichier_export = self._export_automatique()

            self._set_progress(
                100,
                len(self.df),
                total
            )

            if fichier_export:

                self.info_label.setText(
                    f"Extraction terminée avec succès — "
                    f"{len(self.df):,} lignes traitées — "
                    f"Export automatique effectué dans : "
                    f"{DOSSIER_EXPORT}".replace(
                        ",",
                        " "
                    )
                )

            else:

                self.info_label.setText(
                    f"Extraction terminée avec succès — "
                    f"{len(self.df):,} lignes traitées.".replace(
                        ",",
                        " "
                    )
                )

            QApplication.processEvents()

            self._close_progress()

            return True

        except Exception as e:

            self._close_progress()

            QMessageBox.critical(
                self,
                "Erreur Garanties",
                f"Impossible de traiter le fichier :\n\n{e}"
            )

            return False

    # ========================================================
    # LECTURE EXCEL
    # ========================================================

    def _lire_excel(
        self,
        fichier
    ):

        xls = pd.ExcelFile(
            fichier
        )

        feuilles = xls.sheet_names

        if not feuilles:
            return pd.DataFrame()

        if len(feuilles) == 1:

            feuille = feuilles[0]

        else:

            feuille, ok = QInputDialog.getItem(
                self,
                "Choisir la feuille",
                "Quelle feuille contient les Garanties ?",
                feuilles,
                0,
                False
            )

            if not ok:
                return None

        df = pd.read_excel(
            fichier,
            sheet_name=feuille,
            dtype=object
        )

        df.columns = [
            str(c).strip()
            for c in df.columns
        ]

        return df

    # ========================================================
    # LECTURE CSV
    # ========================================================

    def _lire_csv(
        self,
        fichier
    ):

        try:

            df = pd.read_csv(
                fichier,
                sep=";",
                dtype=object,
                encoding="utf-8-sig",
                keep_default_na=False,
                on_bad_lines="skip"
            )

        except UnicodeDecodeError:

            df = pd.read_csv(
                fichier,
                sep=";",
                dtype=object,
                encoding="latin1",
                keep_default_na=False,
                on_bad_lines="skip"
            )

        df.columns = [
            str(c).strip()
            for c in df.columns
        ]

        return df

    # ========================================================
    # PREPARATION BASE
    # ========================================================

    def _preparer_base(
        self,
        df
    ):

        df = df.copy()

        df = df.dropna(
            how="all"
        ).reset_index(
            drop=True
        )

        df.columns = [
            str(c).strip()
            for c in df.columns
        ]

        colonnes = {

            "code_agence": self._find_column(
                df,
                [
                    "Code Agence",
                    "CodeAgence",
                    "Agence"
                ]
            ),

            "id": self._find_column(
                df,
                [
                    "Id",
                    "ID",
                    "Identifiant"
                ]
            ),

            "relation": self._find_column(
                df,
                [
                    "Relation"
                ]
            ),

            "dossier": self._find_column(
                df,
                [
                    "N° Dossier",
                    "N Dossier",
                    "No Dossier",
                    "Numero Dossier",
                    "Numéro Dossier"
                ]
            ),

            "garantie_num": self._find_column(
                df,
                [
                    "N° Garantie",
                    "N Garantie",
                    "No Garantie",
                    "Numero Garantie",
                    "Numéro Garantie"
                ]
            ),

            "code": self._find_column(
                df,
                [
                    "Code Garantie"
                ]
            ),

            "garantie": self._find_column(
                df,
                [
                    "Garantie",
                    "Libellé Garantie",
                    "Libelle Garantie"
                ]
            ),

            "montant": self._find_column(
                df,
                [
                    "Montant",
                    "Montant Garantie"
                ]
            ),

            "statut": self._find_column(
                df,
                [
                    "Statut",
                    "STATUS",
                    "Status"
                ]
            ),

            "observation": self._find_column(
                df,
                [
                    "Observation",
                    "Observations"
                ]
            )
        }

        colonnes_obligatoires = [
            "code_agence",
            "id",
            "relation",
            "dossier",
            "garantie_num",
            "code",
            "garantie",
            "montant",
            "statut"
        ]

        # ====================================================
        # STRUCTURE FINALE
        # ====================================================

        if all(
            colonnes[c] is not None
            for c in colonnes_obligatoires
        ):

            result = pd.DataFrame()

            result["Code Agence"] = df[
                colonnes["code_agence"]
            ].apply(
                self._nettoyer_code_agence
            )

            result["Id"] = df[
                colonnes["id"]
            ].apply(
                self._nettoyer_entier
            )

            result["Relation"] = df[
                colonnes["relation"]
            ].apply(
                self._texte
            )

            result["N° Dossier"] = df[
                colonnes["dossier"]
            ].apply(
                self._nettoyer_numero
            )

            result["N° Garantie"] = df[
                colonnes["garantie_num"]
            ].apply(
                self._nettoyer_numero
            )

            result["Code Garantie"] = df[
                colonnes["code"]
            ].apply(
                self._texte
            )

            result["Garantie"] = df[
                colonnes["garantie"]
            ].apply(
                self._texte
            )

            result["Montant"] = df[
                colonnes["montant"]
            ].apply(
                self._nettoyer_montant
            )

            result["Statut"] = df[
                colonnes["statut"]
            ].apply(
                self._texte
            )

            if colonnes["observation"] is not None:

                result["Observation"] = df[
                    colonnes["observation"]
                ].apply(
                    self._texte
                )

            else:

                result["Observation"] = ""

            return result.reset_index(
                drop=True
            )

        # ====================================================
        # FICHIER BRUT
        # ====================================================

        brut = df.copy()

        if len(brut.columns) < 10:

            QMessageBox.warning(
                self,
                "Structure du fichier",
                "La structure du fichier Garanties "
                "ne correspond pas à celle attendue."
            )

            return pd.DataFrame()

        brut.drop(
            columns=[
                brut.columns[4]
            ],
            inplace=True
        )

        if len(brut.columns) > 8:

            brut.drop(
                columns=[
                    brut.columns[8]
                ],
                inplace=True
            )

        if len(brut.columns) > 9:

            brut.drop(
                columns=[
                    brut.columns[9]
                ],
                inplace=True
            )

        noms_finaux = [
            "Code Agence",
            "Id",
            "Relation",
            "N° Dossier",
            "N° Garantie",
            "Code Garantie",
            "Garantie",
            "Montant",
            "Statut",
            "Observation"
        ]

        while len(brut.columns) < 10:

            brut[
                f"_vide_{len(brut.columns)}"
            ] = ""

        brut = brut.iloc[
            :,
            :10
        ]

        brut.columns = noms_finaux

        brut["Code Agence"] = brut[
            "Code Agence"
        ].apply(
            self._nettoyer_code_agence
        )

        brut["Id"] = brut[
            "Id"
        ].apply(
            self._nettoyer_entier
        )

        brut["Relation"] = brut[
            "Relation"
        ].apply(
            self._texte
        )

        brut["N° Dossier"] = brut[
            "N° Dossier"
        ].apply(
            self._nettoyer_numero
        )

        brut["N° Garantie"] = brut[
            "N° Garantie"
        ].apply(
            self._nettoyer_numero
        )

        brut["Code Garantie"] = brut[
            "Code Garantie"
        ].apply(
            self._texte
        )

        brut["Garantie"] = brut[
            "Garantie"
        ].apply(
            self._texte
        )

        brut["Montant"] = brut[
            "Montant"
        ].apply(
            self._nettoyer_montant
        )

        brut["Statut"] = brut[
            "Statut"
        ].apply(
            self._texte
        )

        brut["Observation"] = brut[
            "Observation"
        ].apply(
            self._texte
        )

        return brut.reset_index(
            drop=True
        )

    # ========================================================
    # TEXTE
    # ========================================================

    def _texte(
        self,
        valeur
    ):

        if valeur is None:
            return ""

        try:

            if pd.isna(valeur):
                return ""

        except Exception:
            pass

        return str(
            valeur
        ).strip()

    # ========================================================
    # NUMERO
    # ========================================================

    def _nettoyer_numero(
        self,
        valeur
    ):

        if valeur is None:
            return ""

        try:

            if pd.isna(valeur):
                return ""

        except Exception:
            pass

        texte = str(
            valeur
        ).strip()

        if texte == "":
            return ""

        if texte.endswith(".0"):

            partie = texte[:-2]

            if partie.isdigit():
                texte = partie

        return texte

    # ========================================================
    # CODE AGENCE
    # ========================================================

    def _nettoyer_code_agence(
        self,
        valeur
    ):

        if valeur is None:
            return ""

        try:

            if pd.isna(valeur):
                return ""

        except Exception:
            pass

        texte = str(
            valeur
        ).strip()

        if texte == "":
            return ""

        if texte.endswith(".0"):

            partie = texte[:-2]

            if partie.isdigit():
                texte = partie

        if texte.isdigit():

            return texte.zfill(
                4
            )

        return texte

    # ========================================================
    # ID
    # ========================================================

    def _nettoyer_entier(
        self,
        valeur
    ):

        if valeur is None:
            return ""

        try:

            if pd.isna(valeur):
                return ""

        except Exception:
            pass

        texte = str(
            valeur
        ).strip()

        if texte == "":
            return ""

        if texte.endswith(".0"):

            partie = texte[:-2]

            if partie.isdigit():
                return partie

        return texte

    # ========================================================
    # MONTANT
    # ========================================================

    def _nettoyer_montant(
        self,
        valeur
    ):

        if valeur is None:
            return ""

        try:

            if pd.isna(valeur):
                return ""

        except Exception:
            pass

        texte = str(
            valeur
        ).strip()

        if texte == "":
            return ""

        texte_sans_espace = texte.replace(
            " ",
            ""
        )

        if texte_sans_espace in (
            "0",
            "0.0",
            "0,0"
        ):

            return "0"

        if (
            "," in texte_sans_espace
            and "." not in texte_sans_espace
        ):

            texte_sans_espace = (
                texte_sans_espace.replace(
                    ",",
                    "."
                )
            )

        try:

            nombre = float(
                texte_sans_espace
            )

        except (
            ValueError,
            TypeError
        ):

            return texte

        if nombre.is_integer():

            return f"{int(nombre):,}".replace(
                ",",
                " "
            )

        return f"{nombre:,.2f}".replace(
            ",",
            " "
        )

    # ========================================================
    # FIND COLUMN
    # ========================================================

    def _find_column(
        self,
        df,
        candidats
    ):

        colonnes_normalisees = {}

        for col in df.columns:

            nom = self._normaliser_nom(
                col
            )

            colonnes_normalisees[
                nom
            ] = col

        for candidat in candidats:

            nom = self._normaliser_nom(
                candidat
            )

            if nom in colonnes_normalisees:

                return colonnes_normalisees[
                    nom
                ]

        return None

    # ========================================================
    # NORMALISATION
    # ========================================================

    def _normaliser_nom(
        self,
        nom
    ):

        return (
            str(nom)
            .strip()
            .lower()
            .replace("é", "e")
            .replace("è", "e")
            .replace("ê", "e")
            .replace("ë", "e")
            .replace("à", "a")
            .replace("â", "a")
            .replace("î", "i")
            .replace("ï", "i")
            .replace("ô", "o")
            .replace("ù", "u")
            .replace("û", "u")
            .replace("ç", "c")
        )

    # ========================================================
    # AFFICHAGE
    # ========================================================

    def _afficher_resultats(
        self
    ):

        self._remplir_table(
            self.base_model,
            self.df
        )

        if "Statut" not in self.df.columns:
            return

        statut = (
            self.df["Statut"]
            .astype(str)
            .str.strip()
            .str.upper()
        )

        constituee = self.df[
            statut.isin(
                [
                    "CONSTITUEE",
                    "CONSTITUÉE"
                ]
            )
        ]

        self._remplir_table(
            self.constituee_model,
            constituee
        )

        incomplete = self.df[
            statut.isin(
                [
                    "INCOMPLETE",
                    "INCOMPLÈTE"
                ]
            )
        ]

        self._remplir_table(
            self.incomplete_model,
            incomplete
        )

        retenue = self.df[
            statut.isin(
                [
                    "NON CONSTITUE",
                    "NON CONSTITUÉ",
                    "NON CONSTITUEE",
                    "NON CONSTITUÉE",
                    "RETENUE"
                ]
            )
        ]

        self._remplir_table(
            self.retenue_model,
            retenue
        )

        echue = self.df[
            statut == "ECHUE"
        ]

        self._remplir_table(
            self.echue_model,
            echue
        )

    # ========================================================
    # REMPLIR MODELE
    # ========================================================

    def _remplir_table(
        self,
        model,
        df
    ):

        if df is None:
            df = pd.DataFrame()

        model.set_dataframe(
            df
        )

        table = self._table_for_model(
            model
        )

        if table is None:
            return

        self._configurer_largeurs(
            table,
            df
        )

    # ========================================================
    # TABLE POUR MODELE
    # ========================================================

    def _table_for_model(
        self,
        model
    ):

        if model is self.base_model:
            return self.base_table

        if model is self.constituee_model:
            return self.constituee_table

        if model is self.incomplete_model:
            return self.incomplete_table

        if model is self.retenue_model:
            return self.retenue_table

        if model is self.echue_model:
            return self.echue_table

        return None

    # ========================================================
    # LARGEURS TABLEAU
    # ========================================================

    def _configurer_largeurs(
        self,
        table,
        df
    ):

        if df is None or df.empty:
            return

        header = table.horizontalHeader()

        for i, colonne in enumerate(df.columns):

            nom = str(
                colonne
            ).lower()

            largeur = 120

            if "observation" in nom:

                largeur = 450

            elif nom == "relation":

                largeur = 330

            elif nom == "garantie":

                largeur = 220

            elif "code agence" in nom:

                largeur = 120

            elif "dossier" in nom:

                largeur = 140

            elif "n° garantie" in nom:

                largeur = 140

            elif "montant" in nom:

                largeur = 150

            elif nom == "id":

                largeur = 100

            elif "code garantie" in nom:

                largeur = 140

            header.resizeSection(
                i,
                largeur
            )

    # ========================================================
    # EXPORT AUTOMATIQUE
    # ========================================================

    def _export_automatique(
        self
    ):

        if not self.fichier_source:
            return None

        if self.df is None or self.df.empty:
            return None

        # ====================================================
        # DOSSIER FIXE
        # ====================================================

        dossier_export = DOSSIER_EXPORT

        try:

            os.makedirs(
                dossier_export,
                exist_ok=True
            )

        except Exception as e:

            QMessageBox.warning(
                self,
                "Export Garanties",
                "Impossible de créer le dossier d'export :\n\n"
                f"{dossier_export}\n\n"
                f"Détail : {e}"
            )

            return None

        # ====================================================
        # NOM DU FICHIER SOURCE
        # ====================================================

        nom_source = os.path.splitext(
            os.path.basename(
                self.fichier_source
            )
        )[0]

        fichier_export = os.path.join(
            dossier_export,
            f"{nom_source}_Garanties_Export.xlsx"
        )

        try:

            # =================================================
            # TABLEAUX
            # =================================================

            tableaux = {

                "Base Garanties":
                    self.base_model._df.copy(),

                "Constituée":
                    self.constituee_model._df.copy(),

                "Incomplète":
                    self.incomplete_model._df.copy(),

                "Retenue NC":
                    self.retenue_model._df.copy(),

                "Echue":
                    self.echue_model._df.copy(),
            }

            # =================================================
            # COULEURS DES ONGLETS
            # =================================================

            couleurs_onglets = {

                "Base Garanties": "008000",

                "Constituée": "FFFF00",

                "Incomplète": "F4B183",

                "Retenue NC": "FF0000",

                "Echue": "0070C0",
            }

            # =================================================
            # IMPORT OPENPYXL
            # =================================================

            from openpyxl.styles import (
                Font,
                PatternFill,
                Alignment,
                Border,
                Side
            )

            # =================================================
            # CREATION EXCEL
            # =================================================

            with pd.ExcelWriter(
                fichier_export,
                engine="openpyxl"
            ) as writer:

                for nom_feuille, dataframe in tableaux.items():

                    if dataframe is None:

                        dataframe = pd.DataFrame()

                    # =========================================
                    # IMPORTANT :
                    # OBSERVATION = COLONNE J
                    # =========================================

                    colonnes_standard = [
                        "Code Agence",
                        "Id",
                        "Relation",
                        "N° Dossier",
                        "N° Garantie",
                        "Code Garantie",
                        "Garantie",
                        "Montant",
                        "Statut",
                        "Observation"
                    ]

                    for colonne in colonnes_standard:

                        if colonne not in dataframe.columns:

                            dataframe[colonne] = ""

                    dataframe = dataframe[
                        colonnes_standard
                    ].copy()

                    # =========================================
                    # ECRITURE
                    # =========================================

                    dataframe.to_excel(
                        writer,
                        sheet_name=nom_feuille,
                        index=False
                    )

                    worksheet = writer.sheets[
                        nom_feuille
                    ]

                    # =========================================
                    # COULEUR ONGLET
                    # =========================================

                    couleur_onglet = (
                        couleurs_onglets.get(
                            nom_feuille
                        )
                    )

                    if couleur_onglet:

                        worksheet.sheet_properties.tabColor = (
                            couleur_onglet
                        )

                    # =========================================
                    # STYLE EN-TETE
                    # =========================================

                    rouge_entete = "C00000"

                    bordure_blanche = Side(
                        style="thin",
                        color="FFFFFF"
                    )

                    for cellule in worksheet[1]:

                        cellule.fill = PatternFill(
                            fill_type="solid",
                            fgColor=rouge_entete
                        )

                        cellule.font = Font(
                            name="Arial",
                            size=11,
                            bold=True,
                            color="FFFFFF"
                        )

                        cellule.alignment = Alignment(
                            horizontal="center",
                            vertical="center"
                        )

                        cellule.border = Border(
                            left=bordure_blanche,
                            right=bordure_blanche,
                            top=bordure_blanche,
                            bottom=bordure_blanche
                        )

                    # =========================================
                    # HAUTEUR EN-TETE
                    # =========================================

                    worksheet.row_dimensions[
                        1
                    ].height = 25

                    # =========================================
                    # FIGER PREMIERE LIGNE
                    # =========================================

                    worksheet.freeze_panes = "A2"

                    # =========================================
                    # FILTRE
                    # =========================================

                    if worksheet.max_row > 1:

                        worksheet.auto_filter.ref = (
                            worksheet.dimensions
                        )

                    # =========================================
                    # LARGEURS
                    # =========================================

                    for column_cells in worksheet.columns:

                        longueur_max = 0

                        for cell in column_cells:

                            try:

                                valeur = str(
                                    cell.value
                                    if cell.value is not None
                                    else ""
                                )

                                longueur_max = max(
                                    longueur_max,
                                    len(valeur)
                                )

                            except Exception:
                                pass

                        largeur = min(
                            max(
                                longueur_max + 2,
                                10
                            ),
                            50
                        )

                        lettre_colonne = (
                            column_cells[0].column_letter
                        )

                        # =====================================
                        # OBSERVATION J -> Z
                        # =====================================

                        if lettre_colonne == "J":

                            largeur = 25

                        worksheet.column_dimensions[
                            lettre_colonne
                        ].width = largeur

                    # =========================================
                    # OBSERVATION DE J A Z
                    # =========================================
                    #
                    # La colonne J contient Observation.
                    # On lui donne une largeur de départ de 25.
                    #
                    # Si tu veux réellement fusionner J:Z,
                    # ce serait une autre logique.
                    #
                    # Ici, les données restent intactes.
                    # =========================================

                    worksheet.column_dimensions[
                        "J"
                    ].width = 25

            # =================================================
            # VERIFICATION
            # =================================================

            if os.path.exists(
                fichier_export
            ):

                return fichier_export

            return None

        except PermissionError:

            QMessageBox.warning(
                self,
                "Export Garanties",
                "Impossible d'effectuer l'export automatique.\n\n"
                "Le fichier Excel est peut-être déjà ouvert."
            )

            return None

        except Exception as e:

            QMessageBox.warning(
                self,
                "Export Garanties",
                "Les données ont été traitées, mais "
                "l'export automatique a échoué :\n\n"
                f"{e}"
            )

            return None

    # ========================================================
    # LANCER
    # ========================================================

    def lancer(
        self
    ):

        return self.import_csv()