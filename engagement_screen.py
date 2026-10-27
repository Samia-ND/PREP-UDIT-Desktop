# ============================================================
# engagement_screen.py
# État des Engagements
# ============================================================

import sys
import os
from pathlib import Path
from datetime import datetime

import pandas as pd

from PyQt6.QtWidgets import (
    QApplication,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QPushButton,
    QLabel,
    QTableView,
    QFileDialog,
    QFrame,
    QHeaderView,
    QMessageBox,
    QTabWidget,
    QInputDialog,
    QDialog,
    QProgressBar,
)

from PyQt6.QtCore import (
    Qt,
    QThread,
    pyqtSignal,
    QAbstractTableModel,
    QTimer,
)

from PyQt6.QtGui import QFont

# ============================================================
# OPENPYXL
# ============================================================

from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)

from openpyxl.utils import get_column_letter

# ============================================================
# MODULES
# ============================================================

sys.path.insert(
    0,
    "."
)

from modules.module_base import (
    ResultatAudit,
    ThemeResultat,
)

from modules.engagement_module import (
    EngagementModule,
)

from fiche_recap_widget import (
    FicheRecapWidget,
)

# ============================================================
# COULEURS
# ============================================================

RED = "#c0392b"
ORANGE = "#E67E00"
GREEN = "#2d5a2d"
DARK = "#2d2d2d"
WHITE = "#ffffff"
SALMON_BG = "#f6c9a8"
BEIGE = "#ead6c8"

# ============================================================
# WORKER QTHREAD
# ============================================================

class EngagementWorker(QThread):

    finished = pyqtSignal(object)
    error = pyqtSignal(str)

    def __init__(
        self,
        module,
        df,
        parent=None
    ):

        super().__init__(parent)

        self.module = module
        self.df = df

    def run(self):

        try:

            # =================================================
            # IMPORTANT :
            # On ne touche PAS au traitement métier.
            # =================================================

            resultat = self.module.executer(
                self.df
            )

            self.finished.emit(
                resultat
            )

        except Exception as e:

            self.error.emit(
                str(e)
            )


# ============================================================
# FENÊTRE DE PROGRESSION
# ============================================================

class ProgressDialog(QDialog):

    def __init__(
        self,
        parent=None
    ):

        super().__init__(
            parent
        )

        self.setFixedSize(
            350,
            120
        )

        self.setWindowTitle(
            "Traitement en cours"
        )

        self.setWindowModality(
            Qt.WindowModality.ApplicationModal
        )

        self.setWindowFlags(
            Qt.WindowType.Dialog
            |
            Qt.WindowType.CustomizeWindowHint
            |
            Qt.WindowType.WindowTitleHint
        )

        self.setStyleSheet(
            """
            QDialog {
                background-color: white;
                border: 2px solid #8C4318;
                border-radius: 8px;
            }

            QLabel {
                background: transparent;
                color: #2d2d2d;
            }

            QProgressBar {
                border: 1px solid #999999;
                border-radius: 6px;
                background-color: #eeeeee;
                height: 20px;
                text-align: center;
                font-size: 11px;
                font-weight: bold;
            }

            QProgressBar::chunk {
                background-color: #c0392b;
                border-radius: 5px;
            }
            """
        )

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            18,
            12,
            18,
            12
        )

        layout.setSpacing(
            7
        )

        # ----------------------------------------------------
        # TITRE
        # ----------------------------------------------------

        self.title_label = QLabel(
            "Traitement des engagements en cours..."
        )

        self.title_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.title_label.setFont(
            QFont(
                "Arial",
                11,
                QFont.Weight.Bold
            )
        )

        # ----------------------------------------------------
        # BARRE
        # ----------------------------------------------------

        self.progress_bar = QProgressBar()

        self.progress_bar.setRange(
            0,
            100
        )

        self.progress_bar.setValue(
            0
        )

        self.progress_bar.setTextVisible(
            True
        )

        # ----------------------------------------------------
        # LIGNES
        # ----------------------------------------------------

        self.lines_label = QLabel(
            "> Lignes traitées : 0 / 0"
        )

        self.lines_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.lines_label.setFont(
            QFont(
                "Arial",
                9
            )
        )

        # ----------------------------------------------------
        # MESSAGE
        # ----------------------------------------------------

        self.wait_label = QLabel(
            "Veuillez patienter..."
        )

        self.wait_label.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        self.wait_label.setFont(
            QFont(
                "Arial",
                9,
                QFont.Weight.Bold
            )
        )

        self.wait_label.setStyleSheet(
            """
            color: #555555;
            """
        )

        layout.addWidget(
            self.title_label
        )

        layout.addWidget(
            self.progress_bar
        )

        layout.addWidget(
            self.lines_label
        )

        layout.addWidget(
            self.wait_label
        )

    # ========================================================
    # UPDATE PROGRESSION
    # ========================================================

    def update_progress(
        self,
        percent,
        processed=0,
        total=0
    ):

        percent = max(
            0,
            min(
                100,
                int(percent)
            )
        )

        self.progress_bar.setValue(
            percent
        )

        if total:

            self.lines_label.setText(
                f"> Lignes traitées : "
                f"{processed:,} / {total:,}"
                .replace(
                    ",",
                    " "
                )
            )

        else:

            self.lines_label.setText(
                f"> Lignes traitées : "
                f"{processed:,}"
                .replace(
                    ",",
                    " "
                )
            )


# ============================================================
# MODÈLE TABLEAU OPTIMISÉ
# ============================================================

class DataFrameModel(QAbstractTableModel):

    """
    Modèle Qt optimisé pour afficher un DataFrame.

    Contrairement à QTableWidget, on ne crée pas un
    QTableWidgetItem pour chaque cellule.

    Les données restent dans le DataFrame et Qt demande
    uniquement les valeurs nécessaires à l'affichage.
    """

    def __init__(
        self,
        dataframe=None,
        parent=None
    ):

        super().__init__(
            parent
        )

        if dataframe is None:

            self._df = pd.DataFrame()

        else:

            self._df = dataframe

    # ========================================================
    # NOMBRE DE LIGNES
    # ========================================================

    def rowCount(
        self,
        parent=None
    ):

        if parent is not None and parent.isValid():

            return 0

        return len(
            self._df
        )

    # ========================================================
    # NOMBRE DE COLONNES
    # ========================================================

    def columnCount(
        self,
        parent=None
    ):

        if parent is not None and parent.isValid():

            return 0

        return len(
            self._df.columns
        )

    # ========================================================
    # DONNÉE CELLULE
    # ========================================================

    def data(
        self,
        index,
        role=Qt.ItemDataRole.DisplayRole
    ):

        if not index.isValid():

            return None

        if role == Qt.ItemDataRole.DisplayRole:

            try:

                value = self._df.iat[
                    index.row(),
                    index.column()
                ]

            except Exception:

                return ""

            if pd.isna(
                value
            ):

                return ""

            return str(
                value
            )

        # ----------------------------------------------------
        # ALIGNEMENT
        # ----------------------------------------------------

        if role == Qt.ItemDataRole.TextAlignmentRole:

            return (
                Qt.AlignmentFlag.AlignLeft
                |
                Qt.AlignmentFlag.AlignVCenter
            )

        return None

    # ========================================================
    # EN-TÊTES
    # ========================================================

    def headerData(
        self,
        section,
        orientation,
        role=Qt.ItemDataRole.DisplayRole
    ):

        if role != Qt.ItemDataRole.DisplayRole:

            return None

        if orientation == Qt.Orientation.Horizontal:

            if (
                0 <= section
                < len(self._df.columns)
            ):

                return str(
                    self._df.columns[section]
                )

        elif orientation == Qt.Orientation.Vertical:

            return str(
                section + 1
            )

        return None


# ============================================================
# CARTE STATISTIQUE
# ============================================================

class StatCard(QFrame):

    def __init__(
        self,
        label_text
    ):

        super().__init__()

        self.setStyleSheet(
            f"""
            QFrame {{
                background-color: {WHITE};
                border: 1px solid #e0b090;
                border-radius: 8px;
            }}
            """
        )

        layout = QVBoxLayout(
            self
        )

        self.value_label = QLabel(
            "0"
        )

        self.value_label.setStyleSheet(
            f"""
            color: {RED};
            font-size: 20px;
            font-weight: bold;
            border: none;
            """
        )

        self.desc_label = QLabel(
            label_text
        )

        self.desc_label.setStyleSheet(
            f"""
            color: {GREEN};
            font-size: 11px;
            border: none;
            """
        )

        layout.addWidget(
            self.value_label
        )

        layout.addWidget(
            self.desc_label
        )

    def set_value(
        self,
        value
    ):

        self.value_label.setText(
            str(value)
        )


# ============================================================
# ONGLET THEME OPTIMISÉ
# ============================================================

class ThemeTab(QWidget):

    def __init__(
        self,
        theme: ThemeResultat
    ):

        super().__init__()

        self.full_df = theme.lignes

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            10,
            10,
            10,
            10
        )

        # ====================================================
        # TABLEAU
        # ====================================================

        self.table = QTableView()

        # ====================================================
        # STYLE TABLEAU
        # ====================================================

        self.table.setStyleSheet(
            f"""
            QTableView {{
                background-color: white;
                gridline-color: #eeeeee;
                font-size: 11px;
                border: 1px solid #999999;
                alternate-background-color: #fafafa;

                selection-background-color: {BEIGE};
                selection-color: black;
            }}

            QTableView::item {{
                background-color: white;
                color: black;
            }}

            QTableView::item:selected {{
                background-color: {BEIGE};
                color: black;
            }}

            QTableView::item:focus {{
                background-color: {BEIGE};
                color: black;
            }}

            QHeaderView::section {{
                background-color: {RED};
                color: white;
                padding: 4px;
                font-weight: bold;
                border: none;
            }}
            """
        )

        # ----------------------------------------------------
        # MODÈLE
        # ----------------------------------------------------

        self.model = DataFrameModel(
            self.full_df,
            self
        )

        self.table.setModel(
            self.model
        )

        # ----------------------------------------------------
        # CONFIGURATION HORIZONTAL
        # ----------------------------------------------------

        horizontal_header = (
            self.table.horizontalHeader()
        )

        horizontal_header.setSectionResizeMode(
            QHeaderView.ResizeMode.Interactive
        )

        horizontal_header.setStretchLastSection(
            False
        )

        # ----------------------------------------------------
        # CONFIGURATION VERTICAL
        # ----------------------------------------------------

        vertical_header = (
            self.table.verticalHeader()
        )

        vertical_header.setDefaultSectionSize(
            25
        )

        # ----------------------------------------------------
        # PERFORMANCE
        # ----------------------------------------------------

        self.table.setAlternatingRowColors(
            True
        )

        self.table.setSelectionBehavior(
            QTableView.SelectionBehavior.SelectRows
        )

        self.table.setEditTriggers(
            QTableView.EditTrigger.NoEditTriggers
        )

        self.table.setSortingEnabled(
            False
        )

        horizontal_header.setDefaultSectionSize(
            120
        )

        # ----------------------------------------------------
        # AJOUT
        # ----------------------------------------------------

        layout.addWidget(
            self.table
        )


# ============================================================
# ÉCRAN ÉTAT DES ENGAGEMENTS
# ============================================================

class EngagementScreen(QWidget):

    def __init__(
        self,
        on_back=None
    ):

        super().__init__()

        self.on_back = on_back

        # ====================================================
        # CALLBACK VERS MAIN.PY
        # ====================================================

        self.on_results_ready = None

        self.module = EngagementModule()

        self.progress_dialog = None

        self.worker = None

        # ====================================================
        # DONNÉES DU RÉSULTAT
        # ====================================================

        self.current_resultat = None

        self.current_source = ""

        self.export_path = None

        self.setStyleSheet(
            f"""
            QWidget {{
                background-color: {SALMON_BG};
                color: {DARK};
            }}

            QTabWidget::pane {{
                border: 1px solid #999999;
                background: white;
            }}

            QTabBar::tab {{
                background: {BEIGE};
                padding: 9px 18px;
                border: 1px solid #999999;
                font-weight: bold;
                color: {DARK};
            }}

            QTabBar::tab:selected {{
                background: white;
                color: {RED};
            }}
            """
        )

        self._build_ui()

    # ========================================================
    # INTERFACE
    # ========================================================

    def _build_ui(
        self
    ):

        layout = QVBoxLayout(
            self
        )

        layout.setContentsMargins(
            20,
            15,
            20,
            20
        )

        layout.setSpacing(
            12
        )

        # ====================================================
        # HEADER
        # ====================================================

        header = QFrame()

        header.setStyleSheet(
            """
            QFrame {
                background-color: white;
                border: 2px solid #8C4318;
                border-radius: 8px;
            }
            """
        )

        header_layout = QHBoxLayout(
            header
        )

        header_layout.setContentsMargins(
            15,
            10,
            15,
            10
        )

        # ----------------------------------------------------
        # RETOUR
        # ----------------------------------------------------

        self.btn_back = QPushButton(
            "← Retour au Menu"
        )

        self.btn_back.setFixedSize(
            142,
            38
        )

        self.btn_back.setCursor(
            Qt.CursorShape.PointingHandCursor
        )

        self.btn_back.setStyleSheet(
            """
            QPushButton {
                background-color: #E67E00;
                color: white;
                border: none;
                border-radius: 6px;
                padding: 5px 12px;
                font-weight: bold;
                font-size: 12px;
            }

            QPushButton:hover {
                background-color: #D66F00;
            }

            QPushButton:pressed {
                background-color: #BF6000;
            }
            """
        )

        self.btn_back.clicked.connect(
            self._go_back
        )

        # ----------------------------------------------------
        # TITRE
        # ----------------------------------------------------

        title = QLabel(
            "ÉTAT DES ENGAGEMENTS"
        )

        title.setFont(
            QFont(
                "Georgia",
                28,
                QFont.Weight.Bold
            )
        )

        title.setStyleSheet(
            f"""
            QLabel {{
                color: {RED};
                border: none;
                background: transparent;
            }}
            """
        )

        title.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        # ----------------------------------------------------
        # ESPACE DROITE
        # ----------------------------------------------------

        right_space = QLabel()

        right_space.setFixedWidth(
            142
        )

        right_space.setStyleSheet(
            """
            QLabel {
                background: transparent;
                border: none;
            }
            """
        )

        header_layout.addWidget(
            self.btn_back
        )

        header_layout.addStretch()

        header_layout.addWidget(
            title
        )

        header_layout.addStretch()

        header_layout.addWidget(
            right_space
        )

        layout.addWidget(
            header
        )

        # ====================================================
        # ONGLETS
        # ====================================================

        self.tabs = QTabWidget()

        self.tabs.setStyleSheet(
            f"""
            QTabWidget::pane {{
                border: 1px solid #999999;
                background: white;
            }}

            QTabBar::tab {{
                background: {BEIGE};
                padding: 9px 18px;
                border: 1px solid #999999;
                font-weight: bold;
                color: {DARK};
            }}

            QTabBar::tab:selected {{
                background: white;
                color: {RED};
            }}
            """
        )

        layout.addWidget(
            self.tabs,
            1
        )

    # ========================================================
    # RETOUR
    # ========================================================

    def _go_back(
        self
    ):

        if self.worker is not None:

            if self.worker.isRunning():

                return

        if self.on_back:

            self.on_back()

    # ========================================================
    # IMPORT
    # ========================================================

    def import_csv(
        self
    ):

        path, _ = QFileDialog.getOpenFileName(
            self,
            "Choisir un état des engagements",
            "",
            "Fichiers supportés (*.xlsx *.xls *.xlsm *.csv *.txt);;"
            "Fichiers Excel (*.xlsx *.xls *.xlsm);;"
            "Fichiers CSV (*.csv *.txt);;"
            "Tous les fichiers (*)"
        )

        if not path:

            return False

        try:

            if path.lower().endswith(
                (
                    ".xlsx",
                    ".xls",
                    ".xlsm"
                )
            ):

                df = (
                    self._lire_excel_avec_choix_feuille(
                        path
                    )
                )

                if df is None:

                    return False

            else:

                df = pd.read_csv(
                    path,
                    sep=";",
                    dtype=str,
                    encoding="utf-8",
                    on_bad_lines="skip"
                )

                df.columns = (
                    df.columns
                    .astype(str)
                    .str.strip()
                )

            return self.process_dataframe(
                df,
                source=path
            )

        except UnicodeDecodeError:

            try:

                df = pd.read_csv(
                    path,
                    sep=";",
                    dtype=str,
                    encoding="latin1",
                    on_bad_lines="skip"
                )

                df.columns = (
                    df.columns
                    .astype(str)
                    .str.strip()
                )

                return self.process_dataframe(
                    df,
                    source=path
                )

            except Exception as e:

                QMessageBox.critical(
                    self,
                    "Erreur de lecture",
                    f"Impossible de lire le fichier :\n{e}"
                )

                return False

        except Exception as e:

            QMessageBox.critical(
                self,
                "Erreur de lecture",
                f"Impossible de lire le fichier :\n{e}"
            )

            return False

    # ========================================================
    # LECTURE EXCEL
    # ========================================================

    def _lire_excel_avec_choix_feuille(
        self,
        path
    ):

        xls = pd.ExcelFile(
            path
        )

        noms_feuilles = xls.sheet_names

        if not noms_feuilles:

            QMessageBox.warning(
                self,
                "Fichier vide",
                "Aucune feuille Excel trouvée."
            )

            return None

        if len(noms_feuilles) == 1:

            feuille = noms_feuilles[0]

        else:

            feuille, ok = QInputDialog.getItem(
                self,
                "Choisir la feuille",
                "Quelle feuille contient l'état des engagements ?",
                noms_feuilles,
                0,
                False
            )

            if not ok:

                return None

        df = pd.read_excel(
            path,
            sheet_name=feuille,
            dtype=str
        )

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
        )

        return df

    # ========================================================
    # DONNÉES EXEMPLE
    # ========================================================

    def load_sample(
        self
    ):

        data = {

            "N° Client": [
                "2623",
                "2623",
                "2623",
                "2623",
                "175498",
                "1087320"
            ],

            "N° Dossier": [
                "00010048",
                "00010048",
                "00000000",
                "00010048",
                "00034434",
                "00017001"
            ],

            "Code Agence": [
                "0908",
                "0908",
                "0908",
                "0908",
                "0908",
                "0908"
            ],

            "Libellé Agence": [
                "C.D'AFF CASA ZENATA",
                "C.D'AFF CASA ZENATA",
                "C.D'AFF CASA ZENATA",
                "C.D'AFF CASA ZENATA",
                "C.D'AFF CASA ZENATA",
                "C.D'AFF CASA ZENATA"
            ],

            "N° Compte": [
                "000908X000000135",
                "000908N000126750",
                "000908N000126750",
                "000908R000001164",
                "000908R000000368",
                "000908B000126792"
            ],

            "Réf Produit": [
                "112",
                "923",
                "812",
                "112",
                "112",
                "923"
            ],

            "Montant Auto": [
                0,
                0,
                0,
                0,
                0,
                0
            ],

            "Montant Util": [
                247366,
                163522424,
                1,
                66,
                26085,
                922218
            ]
        }

        df = pd.DataFrame(
            data
        )

        return self.process_dataframe(
            df,
            source="Données d'exemple"
        )

    # ========================================================
    # TRAITEMENT
    # ========================================================

    def process_dataframe(
        self,
        df,
        source=""
    ):

        if df is None or df.empty:

            QMessageBox.warning(
                self,
                "Fichier vide",
                "Le fichier importé ne contient aucune donnée."
            )

            return False

        df = df.copy()

        # ====================================================
        # MÉMORISER LA SOURCE
        # ====================================================

        self.current_source = source

        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
        )

        # ====================================================
        # VALIDATION
        # ====================================================

        missing = self.module.valider(
            df
        )

        if missing:

            QMessageBox.warning(
                self,
                "Colonnes manquantes",
                "Le fichier ne contient pas les colonnes "
                "nécessaires au traitement :\n\n"
                +
                "\n".join(
                    f"• {col}"
                    for col in missing
                )
            )

            return False

        # ====================================================
        # FENÊTRE DE PROGRESSION
        # ====================================================

        self.progress_dialog = ProgressDialog(
            self
        )

        total = len(
            df
        )

        self.progress_dialog.update_progress(
            0,
            0,
            total
        )

        self.progress_dialog.title_label.setText(
            "Traitement des engagements en cours..."
        )

        self.progress_dialog.wait_label.setText(
            "Analyse des données..."
        )

        self.progress_dialog.show()

        # ====================================================
        # IMPORTANT
        # ====================================================

        QApplication.processEvents()

        # ====================================================
        # VERROUILLER LE RETOUR
        # ====================================================

        self.btn_back.setEnabled(
            False
        )

        # ====================================================
        # THREAD
        # ====================================================

        self.worker = EngagementWorker(
            self.module,
            df,
            self
        )

        # ====================================================
        # SIGNAUX
        # ====================================================

        self.worker.finished.connect(
            self._traitement_termine
        )

        self.worker.error.connect(
            self._traitement_erreur
        )

        # ====================================================
        # LANCEMENT
        # ====================================================

        self.worker.start()

        return True

    # ========================================================
    # TRAITEMENT TERMINÉ
    # ========================================================

    def _traitement_termine(
        self,
        resultat
    ):

        # ====================================================
        # NOMBRE TOTAL DE LIGNES
        # ====================================================

        total_lines = 0

        if self.worker:

            total_lines = len(
                self.worker.df
            )

        # ====================================================
        # POPUP
        # ====================================================

        if self.progress_dialog:

            self.progress_dialog.update_progress(
                100,
                total_lines,
                total_lines
            )

            self.progress_dialog.title_label.setText(
                "Préparation des résultats..."
            )

            self.progress_dialog.wait_label.setText(
                "Construction des résultats..."
            )

        # ====================================================
        # SUPPRESSION DES ANCIENS ONGLETS
        # ====================================================

        self.tabs.clear()

        # ====================================================
        # FICHE RÉCAP
        # ====================================================

        if resultat.fiche_recap:

            fiche_recap = FicheRecapWidget(
                resultat.fiche_recap,
                resultat.date_execution
            )

            self.tabs.addTab(
                fiche_recap,
                "Fiche Récap"
            )

        # ====================================================
        # ORDRE DES THÈMES
        # ====================================================

        ordre = [

            "Nature AU",

            "Total Engagement",

            "Dep. Util AU",

            "Base AU Hrs CTX",

            "Contentieux AU",

            "Dossiers Echus AU",

            "Impayés ESC & CMT",

            "Sans Dossier AU",

            "CMT Oxygène",

            "Relance",

            "ESC COM",

            "ACNE",

            "ASMN",

            "ASM CR COMP"
        ]

        # ====================================================
        # INDEX DES THÈMES
        # ====================================================

        themes_dict = {

            theme.nom_theme: theme

            for theme in resultat.themes
        }

        # ====================================================
        # CRÉATION DES THÈMES
        # ====================================================

        for nom_theme in ordre:

            theme = themes_dict.get(
                nom_theme
            )

            if theme is None:

                continue

            theme_tab = ThemeTab(
                theme
            )

            self.tabs.addTab(
                theme_tab,
                f"{theme.nom_theme} ({theme.nb_lignes})"
            )

        # ====================================================
        # MÉMORISER LE RÉSULTAT
        # ====================================================

        self.current_resultat = resultat

        # ====================================================
        # FERMER LA POPUP
        # ====================================================

        if self.progress_dialog:

            self.progress_dialog.title_label.setText(
                "Traitement terminé"
            )

            self.progress_dialog.wait_label.setText(
                "Terminé."
            )

            self.progress_dialog.update_progress(
                100,
                total_lines,
                total_lines
            )

            self.progress_dialog.close()

            self.progress_dialog.deleteLater()

            self.progress_dialog = None

        # ====================================================
        # RÉACTIVER LE BOUTON
        # ====================================================

        self.btn_back.setEnabled(
            True
        )

        # ====================================================
        # NETTOYAGE THREAD
        # ====================================================

        worker = self.worker

        self.worker = None

        if worker:

            worker.deleteLater()

        # ====================================================
        # NAVIGATION
        # ====================================================

        if self.on_results_ready:

            self.on_results_ready(
                self
            )

        # ====================================================
        # EXPORT EXCEL AUTOMATIQUE
        # ====================================================

        QTimer.singleShot(
            300,
            lambda: self._export_resultat_excel(
                resultat
            )
        )

        return True

    # ========================================================
    # EXPORT EXCEL AUTOMATIQUE
    # ========================================================

    def _export_resultat_excel(
        self,
        resultat
    ):

        try:

            # =================================================
            # DOSSIER EXACT DEMANDÉ
            #
            # C:\Users\HP\Desktop\prep@udit\EXTRACTION
            # =================================================

            export_dir = Path(
                r"C:\Users\HP\Desktop\prep@udit\EXTRACTION"
            )

            export_dir.mkdir(
                parents=True,
                exist_ok=True
            )

            # =================================================
            # NOM DU FICHIER
            # =================================================

            date_export = datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )

            filename = (
                f"Etat_Engagements_"
                f"{date_export}.xlsx"
            )

            output_path = (
                export_dir
                /
                filename
            )

            # =================================================
            # CLASSEUR
            # =================================================

            wb = Workbook()

            # Supprimer la feuille par défaut

            default_sheet = wb.active

            wb.remove(
                default_sheet
            )

            # =================================================
            # STYLES
            # =================================================

            red_fill = PatternFill(
                fill_type="solid",
                fgColor="C0392B"
            )

            cream_fill = PatternFill(
                fill_type="solid",
                fgColor="FFF2CC"
            )

            white_font = Font(
                color="FFFFFF",
                bold=True
            )

            bold_font = Font(
                bold=True,
                color="000000"
            )

            normal_font = Font(
                color="000000"
            )

            thin_side = Side(
                style="thin",
                color="999999"
            )

            thin_border = Border(
                left=thin_side,
                right=thin_side,
                top=thin_side,
                bottom=thin_side
            )

            center = Alignment(
                horizontal="center",
                vertical="center"
            )

            right = Alignment(
                horizontal="right",
                vertical="center"
            )

            left = Alignment(
                horizontal="left",
                vertical="center"
            )

            # =================================================
            # FICHE RÉCAP
            # =================================================

            fiche = resultat.fiche_recap

            if fiche:

                ws = wb.create_sheet(
                    "Fiche Récap"
                )

                # =================================================
                # LARGEURS A → I
                # =================================================

                column_widths = {

                    "A": 18,
                    "B": 18,
                    "C": 18,

                    "D": 18,
                    "E": 18,
                    "F": 18,

                    "G": 18,
                    "H": 18,
                    "I": 18,
                }

                for col, width in column_widths.items():

                    ws.column_dimensions[
                        col
                    ].width = width

                # =================================================
                # DATE
                # =================================================

                ws.merge_cells(
                    "G1:H1"
                )

                ws["G1"] = "Date"

                ws["G1"].font = Font(
                    bold=True,
                    italic=True,
                    color="000000"
                )

                ws["G1"].alignment = Alignment(
                    horizontal="right",
                    vertical="center"
                )

                ws["I1"] = (

                    resultat.date_execution.strftime(
                        "%d/%m/%Y"
                    )

                    if resultat.date_execution

                    else ""
                )

                ws["I1"].font = Font(
                    bold=True,
                    italic=True,
                    color="000000"
                )

                ws["I1"].alignment = Alignment(
                    horizontal="left",
                    vertical="center"
                )

                ws.row_dimensions[
                    1
                ].height = 24

                current_row = 3

                # =================================================
                # NATURE DES ENGAGEMENTS
                # =================================================

                if "nature" in fiche:

                    # -------------------------------------------------
                    # TITRE A:I
                    # -------------------------------------------------

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=1,
                        end_row=current_row,
                        end_column=9
                    )

                    cell = ws.cell(
                        current_row,
                        1,
                        "NATURE DES ENGAGEMENTS"
                    )

                    cell.fill = red_fill
                    cell.font = white_font
                    cell.alignment = center
                    cell.border = thin_border

                    for col in range(1, 10):

                        c = ws.cell(
                            current_row,
                            col
                        )

                        c.fill = red_fill
                        c.border = thin_border

                    ws.row_dimensions[
                        current_row
                    ].height = 28

                    current_row += 1

                    # -------------------------------------------------
                    # EN-TÊTES
                    #
                    # A:C = NATURE
                    # D:F = UTILISATIONS
                    # G:I = POURCENTAGE
                    # -------------------------------------------------

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=1,
                        end_row=current_row,
                        end_column=3
                    )

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=4,
                        end_row=current_row,
                        end_column=6
                    )

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=7,
                        end_row=current_row,
                        end_column=9
                    )

                    ws.cell(
                        current_row,
                        1,
                        "NATURE"
                    )

                    ws.cell(
                        current_row,
                        4,
                        "UTILISATIONS"
                    )

                    ws.cell(
                        current_row,
                        7,
                        "POURCENTAGE %"
                    )

                    for col in [1, 4, 7]:

                        cell = ws.cell(
                            current_row,
                            col
                        )

                        cell.fill = red_fill
                        cell.font = white_font
                        cell.alignment = center

                    for col in range(1, 10):

                        cell = ws.cell(
                            current_row,
                            col
                        )

                        cell.fill = red_fill
                        cell.border = thin_border

                    ws.row_dimensions[
                        current_row
                    ].height = 28

                    current_row += 1

                    # -------------------------------------------------
                    # DONNÉES
                    # -------------------------------------------------

                    totals = {
                        "ENGAGEMENTS GLOBALAUX HORS CONTENTIEUX",
                        "ENGAGEMENTS GLOBALAUX"
                    }

                    for label, montant, pct in fiche["nature"]:

                        is_total = (
                            label in totals
                        )

                        # =============================================
                        # NATURE A:C
                        # =============================================

                        ws.merge_cells(
                            start_row=current_row,
                            start_column=1,
                            end_row=current_row,
                            end_column=3
                        )

                        cell_label = ws.cell(
                            current_row,
                            1,
                            str(label)
                        )

                        cell_label.alignment = left
                        cell_label.border = thin_border

                        # =============================================
                        # UTILISATIONS D:F
                        # =============================================

                        ws.merge_cells(
                            start_row=current_row,
                            start_column=4,
                            end_row=current_row,
                            end_column=6
                        )

                        try:

                            montant_value = float(
                                montant
                            )

                        except Exception:

                            montant_value = 0

                        cell_montant = ws.cell(
                            current_row,
                            4,
                            montant_value
                        )

                        cell_montant.number_format = (
                            '#,##0'
                        )

                        cell_montant.alignment = right
                        cell_montant.border = thin_border

                        # =============================================
                        # POURCENTAGE G:I
                        # =============================================

                        ws.merge_cells(
                            start_row=current_row,
                            start_column=7,
                            end_row=current_row,
                            end_column=9
                        )

                        try:

                            pct_value = float(
                                pct
                            )

                        except Exception:

                            pct_value = 0

                        cell_pct = ws.cell(
                            current_row,
                            7,
                            pct_value
                        )

                        cell_pct.number_format = (
                            '0.00"%"'
                        )

                        cell_pct.alignment = right
                        cell_pct.border = thin_border

                        # =============================================
                        # STYLE DE LA LIGNE
                        # =============================================

                        for col in range(1, 10):

                            cell = ws.cell(
                                current_row,
                                col
                            )

                            cell.border = thin_border

                            if is_total:

                                cell.fill = cream_fill
                                cell.font = bold_font

                            else:

                                cell.font = normal_font

                        ws.row_dimensions[
                            current_row
                        ].height = 25

                        current_row += 1

                    current_row += 2

                # =================================================
                # NATURE DES IMPAYÉS
                # =================================================

                if "impayes" in fiche:

                    # -------------------------------------------------
                    # TITRE A:I
                    # -------------------------------------------------

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=1,
                        end_row=current_row,
                        end_column=9
                    )

                    cell = ws.cell(
                        current_row,
                        1,
                        "NATURE DES IMPAYES"
                    )

                    cell.fill = red_fill
                    cell.font = white_font
                    cell.alignment = center
                    cell.border = thin_border

                    for col in range(1, 10):

                        c = ws.cell(
                            current_row,
                            col
                        )

                        c.fill = red_fill
                        c.border = thin_border

                    ws.row_dimensions[
                        current_row
                    ].height = 28

                    current_row += 1

                    # -------------------------------------------------
                    # EN-TÊTES
                    #
                    # A:D = NATURE
                    # E:I = MONTANT
                    # -------------------------------------------------

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=1,
                        end_row=current_row,
                        end_column=4
                    )

                    ws.merge_cells(
                        start_row=current_row,
                        start_column=5,
                        end_row=current_row,
                        end_column=9
                    )

                    ws.cell(
                        current_row,
                        1,
                        "NATURE"
                    )

                    ws.cell(
                        current_row,
                        5,
                        "MONTANT"
                    )

                    for col in [1, 5]:

                        cell = ws.cell(
                            current_row,
                            col
                        )

                        cell.fill = red_fill
                        cell.font = white_font
                        cell.alignment = center

                    for col in range(1, 10):

                        cell = ws.cell(
                            current_row,
                            col
                        )

                        cell.fill = red_fill
                        cell.border = thin_border

                    ws.row_dimensions[
                        current_row
                    ].height = 28

                    current_row += 1

                    # -------------------------------------------------
                    # DONNÉES
                    # -------------------------------------------------

                    for label, montant in fiche["impayes"]:

                        is_total = (
                            label == "TOTAL"
                        )

                        # =============================================
                        # NATURE A:D
                        # =============================================

                        ws.merge_cells(
                            start_row=current_row,
                            start_column=1,
                            end_row=current_row,
                            end_column=4
                        )

                        cell_label = ws.cell(
                            current_row,
                            1,
                            str(label)
                        )

                        if is_total:

                            cell_label.alignment = center
                            cell_label.font = bold_font
                            cell_label.fill = cream_fill

                        else:

                            cell_label.alignment = left
                            cell_label.font = normal_font

                        # =============================================
                        # MONTANT E:I
                        # =============================================

                        ws.merge_cells(
                            start_row=current_row,
                            start_column=5,
                            end_row=current_row,
                            end_column=9
                        )

                        try:

                            montant_value = float(
                                montant
                            )

                        except Exception:

                            montant_value = 0

                        cell_montant = ws.cell(
                            current_row,
                            5,
                            montant_value
                        )

                        cell_montant.number_format = (
                            '#,##0'
                        )

                        cell_montant.alignment = right

                        if is_total:

                            cell_montant.font = bold_font
                            cell_montant.fill = cream_fill

                        else:

                            cell_montant.font = normal_font

                        # =============================================
                        # BORDURES A:I
                        # =============================================

                        for col in range(1, 10):

                            cell = ws.cell(
                                current_row,
                                col
                            )

                            cell.border = thin_border

                            if is_total:

                                cell.fill = cream_fill

                        ws.row_dimensions[
                            current_row
                        ].height = 25

                        current_row += 1

                # =================================================
                # ALIGNEMENT VERTICAL
                # =================================================

                for row in ws.iter_rows(
                    min_row=1,
                    max_row=current_row,
                    min_col=1,
                    max_col=9
                ):

                    for cell in row:

                        current_alignment = cell.alignment

                        cell.alignment = Alignment(
                            horizontal=(
                                current_alignment.horizontal
                                if current_alignment
                                and current_alignment.horizontal
                                else "left"
                            ),
                            vertical="center",
                            wrap_text=False
                        )

                # =================================================
                # FIGER
                # =================================================

                ws.freeze_panes = "A3"

                # =================================================
                # PAS DE COULEUR DE FOND SUR TOUTE LA FEUILLE
                # =================================================
                #
                # On ne met volontairement aucun fill global.
                # Les cellules hors A:I restent blanches.
                #
                # =================================================

            # =================================================
            # ORDRE DES THÈMES
            # =================================================

            ordre = [

                "Nature AU",

                "Total Engagement",

                "Dep. Util AU",

                "Base AU Hrs CTX",

                "Contentieux AU",

                "Dossiers Echus AU",

                "Impayés ESC & CMT",

                "Sans Dossier AU",

                "CMT Oxygène",

                "Relance",

                "ESC COM",

                "ACNE",

                "ASMN",

                "ASM CR COMP"
            ]

            # =================================================
            # INDEX DES THÈMES
            # =================================================

            themes_dict = {

                theme.nom_theme: theme

                for theme in resultat.themes
            }

            # =================================================
            # EXPORT DES THÈMES
            # =================================================

            for nom_theme in ordre:

                theme = themes_dict.get(
                    nom_theme
                )

                if theme is None:

                    continue

                # ------------------------------------------------
                # NOM DE FEUILLE
                #
                # Excel limite le nom d'une feuille à 31 caractères.
                # ------------------------------------------------

                sheet_name = str(
                    theme.nom_theme
                )[:31]

                ws = wb.create_sheet(
                    sheet_name
                )

                df_theme = theme.lignes

                if df_theme is None:

                    continue

                # ------------------------------------------------
                # EN-TÊTES
                # ------------------------------------------------

                for col_index, column_name in enumerate(
                    df_theme.columns,
                    start=1
                ):

                    cell = ws.cell(
                        1,
                        col_index,
                        str(column_name)
                    )

                    cell.fill = red_fill
                    cell.font = white_font
                    cell.alignment = center
                    cell.border = thin_border

                # ------------------------------------------------
                # DONNÉES
                # ------------------------------------------------

                for row_index, row in enumerate(
                    df_theme.itertuples(
                        index=False,
                        name=None
                    ),
                    start=2
                ):

                    for col_index, value in enumerate(
                        row,
                        start=1
                    ):

                        if pd.isna(
                            value
                        ):

                            value = ""

                        cell = ws.cell(
                            row_index,
                            col_index,
                            value
                        )

                        cell.border = thin_border
                        cell.alignment = left

                # ------------------------------------------------
                # LARGEUR DES COLONNES
                # ------------------------------------------------

                for col_index, column_name in enumerate(
                    df_theme.columns,
                    start=1
                ):

                    max_length = len(
                        str(column_name)
                    )

                    sample = df_theme.iloc[
                        :min(
                            len(df_theme),
                            500
                        ),
                        col_index - 1
                    ]

                    for value in sample:

                        if pd.isna(
                            value
                        ):

                            continue

                        length = len(
                            str(value)
                        )

                        if length > max_length:

                            max_length = length

                    width = min(
                        max(
                            max_length + 2,
                            12
                        ),
                        40
                    )

                    ws.column_dimensions[
                        get_column_letter(
                            col_index
                        )
                    ].width = width

                # ------------------------------------------------
                # HAUTEUR DES LIGNES
                # ------------------------------------------------

                ws.row_dimensions[
                    1
                ].height = 25

                for row_number in range(
                    2,
                    len(df_theme) + 2
                ):

                    ws.row_dimensions[
                        row_number
                    ].height = 22

                # ------------------------------------------------
                # FILTRE AUTOMATIQUE
                # ------------------------------------------------

                if len(
                    df_theme.columns
                ) > 0:

                    last_column = get_column_letter(
                        len(
                            df_theme.columns
                        )
                    )

                    last_row = (
                        len(
                            df_theme
                        )
                        +
                        1
                    )

                    ws.auto_filter.ref = (
                        f"A1:{last_column}{last_row}"
                    )

                # ------------------------------------------------
                # FIGER LA PREMIÈRE LIGNE
                # ------------------------------------------------

                ws.freeze_panes = "A2"

            # =================================================
            # SAUVEGARDE
            # =================================================

            wb.save(
                output_path
            )

            self.export_path = str(
                output_path
            )

            print(
                "================================================"
            )

            print(
                "EXPORT EXCEL AUTOMATIQUE"
            )

            print(
                f"Fichier : {output_path}"
            )

            print(
                "================================================"
            )

        except PermissionError:

            QMessageBox.warning(
                self,
                "Export Excel",
                "Impossible d'enregistrer le fichier Excel.\n\n"
                "Le fichier est peut-être déjà ouvert dans Excel.\n\n"
                "Fermez le fichier Excel puis relancez le traitement."
            )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Erreur d'export Excel",
                "Le résultat a bien été affiché, "
                "mais l'export Excel a échoué.\n\n"
                f"{e}"
            )

    # ========================================================
    # ERREUR
    # ========================================================

    def _traitement_erreur(
        self,
        message
    ):

        if self.progress_dialog:

            self.progress_dialog.close()

            self.progress_dialog.deleteLater()

            self.progress_dialog = None

        self.btn_back.setEnabled(
            True
        )

        QMessageBox.critical(
            self,
            "Erreur de traitement",
            "Une erreur est survenue pendant "
            "le traitement :\n\n"
            +
            message
        )

        worker = self.worker

        self.worker = None

        if worker:

            worker.deleteLater()

    # ========================================================
    # MÉTHODE POUR MAIN.PY
    # ========================================================

    def lancer(
        self
    ):

        return self.import_csv()